import json,sys,re,openpyxl
sys.path.insert(0,'evidence/build');from xlstyle import *
from openpyxl.styles import Alignment
Z=json.load(open(sys.argv[1]));PM=json.load(open('final_review/pagemap.json'))
key=lambda u:re.sub(r'^https?://(www\.)?','',u.lower()).rstrip('/').replace('.aspx','')
U={key(v['url']):(v['ar'] or v['title']) for k,v in PM.items() if isinstance(v,dict) and v.get('url')}
EXTRA={'pmu.edu.sa/admission/default':'الصفحة الافتراضية القديمة لقسم القبول','admissions.pmu.edu.sa/student/login/register':'بوابة التقديم – نموذج التسجيل','pmu.edu.sa/admission/mshd_international_students':'ماجستير التنمية البشرية – الطلاب الدوليون','pmu.edu.sa/admission/mba_international_students':'ماجستير إدارة الأعمال – الطلاب الدوليون','pmu.edu.sa/admission/msme_international_students':'ماجستير الهندسة الميكانيكية – الطلاب الدوليون'}
pname=lambda u:U.get(key(u)) or EXTRA.get(key(u)) or u
OWN={'Admissions and relevant programme authority':'مكتب القبول والجهة الأكاديمية المعنية بالبرنامج','Financial Affairs and Admissions':'الشؤون المالية ومكتب القبول','Programme / Registrar authority':'الجهة الأكاديمية / عمادة القبول والتسجيل','Admissions content editor':'محرر محتوى مكتب القبول'}
own=lambda s:next((v for k,v in OWN.items() if s and s.startswith(k)),s)+' (مقترح – يُؤكَّد)'
wb=openpyxl.Workbook();ws=wb.active;ws.title='تصحيحات تحريرية جاهزة'
ws.append(['#','الصفحة','الرابط','الموقع','النص الحالي (حرفيًا)','التصحيح المقترح','حالة التحقق','مرجع الحزمة'])
i=0
for e in sorted(Z['E'],key=lambda e:(not e['_v'],e['change_id'])):
    i+=1
    loc='نموذج التسجيل في بوابة التقديم' if 'register' in e['url'] else 'متن الصفحة'
    ws.append([i,pname(e['url']),e['url'],loc,e['before'],e['after'],'مُتحقَّق منه في المحتوى المنشور' if e['_v'] else 'يحتاج تحققًا يدويًا (يُحمَّل ديناميكيًا)',e['change_id']])
title(ws,'المرفق (أ): تصحيحات تحريرية جاهزة للتنفيذ','بنود حزمة المعالجة الواقعة ضمن نطاق صفحة مكتب القبول فقط؛ لا تغيّر أي قيمة أو سياسة',8)
body(ws,4,[5,30,40,22,36,36,24,11],'مرفق سجل ملاحظات صفحات مكتب القبول',center=(1,7,8))
for row in ws.iter_rows(min_row=5):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if c.column in (1,7,8) else ('left' if c.column in (3,5,6) else 'right'),readingOrder=2 if c.column not in (3,5,6) else 1)
ws.sheet_view.rightToLeft=True
d=wb.create_sheet('قرارات مطلوبة');d.sheet_view.rightToLeft=True
d.append(['#','الموضوع','النص المنشور ذو الصلة','التعارض أو الفجوة','الجهة المالكة','المستند المطلوب للاعتماد','القيمة أو القاعدة المعتمدة','تاريخ السريان','مرجع الاعتماد','مرجع الحزمة'])
DOC='سياسة أو خطة برنامج أو تقويم معتمد ومؤرخ وموقّع يحدد الدفعة والبرنامج وفئة المتقدم وفترة السريان'
for j,x in enumerate(Z['D'],1):
    obs=(x['published_observation'] or '').replace('\n',' | ')
    if re.search(r'\(966\)|Tel\. No',obs): obs='[أسماء موظفين وأرقام هواتف مباشرة – حُذفت حفاظًا على الخصوصية]'
    subj=x['subject'].rstrip('…').strip()
    d.append([j,subj,obs,x['conflict_or_gap'],own(x['owner_proposed']),DOC,'','','',x['decision_id']])
title(d,'المرفق (ب): قرارات مطلوبة من الجهات المالكة قبل تعديل الصياغة','الحقول الثلاثة قبل الأخيرة تُملأ من الجهة المالكة؛ لا تُعدَّل القيمة المنشورة قبل إرفاق مرجع الاعتماد',10)
body(d,4,[5,30,40,44,30,34,22,13,16,11],'مرفق سجل ملاحظات صفحات مكتب القبول',center=(1,10))
for row in d.iter_rows(min_row=5):
    for c in row:
        c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if c.column in (1,10) else ('left' if c.column==3 else 'right'),readingOrder=1 if c.column==3 else 2)
        if c.column in (7,8,9): c.fill=PatternFill('solid',fgColor='FFF9E6')
n=wb.create_sheet('ملاحظات الإعداد');n.sheet_view.rightToLeft=True;n.sheet_view.showGridLines=False
for t in ['المصدر: حزمة «PMU Content Remediation» الشاملة للموقع (530 إجراءً، 101 تصحيح تحريري، 308 مسودات قرار). حكمها الذاتي: غير مكتملة، ولم يُنشر منها شيء.',
          f'اقتُصر هنا على ما يخص صفحة مكتب القبول وصفحاتها وملفاتها المرتبطة: {i} تصحيحًا تحريريًا و{len(Z["D"])} قرارًا. استُبعد ما يخص الكليات والأدلة الأكاديمية والحياة الجامعية والصفحات القديمة غير المرتبطة.',
          'طوبق كل «نص حالي» حرفيًا مع المحتوى المنشور الملتقط في 10–13 أكتوبر 2026: تطابق 44 بندًا، وبقيت 4 تسميات في نموذج التسجيل تُحمَّل ديناميكيًا فوُسمت «يحتاج تحققًا يدويًا».',
          'التصحيحات التحريرية لا تغيّر أي قيمة أو شرط أو رسم؛ أما القرارات فلا يُعدَّل نصها المنشور قبل اعتماد الجهة المالكة وإرفاق المرجع.',
          'حُذفت من البنود أي أسماء أو أرقام هواتف لموظفين، وأُبقيت معرّفات الحزمة (CE/PD) للتتبع فقط.']:
    n.append([t])
title(n,'ملاحظات الإعداد','',1);n.column_dimensions['A'].width=120
for row in n.iter_rows(min_row=4):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='right',readingOrder=2)
wb.save('final_review/مرفق_التصحيحات_والقرارات_المطلوبة.xlsx');print('ok',i,len(Z['D']))
