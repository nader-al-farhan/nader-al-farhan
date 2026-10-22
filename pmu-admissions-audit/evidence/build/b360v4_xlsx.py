# Master Register (Directive 2.1 §23). Usage: b360v4_xlsx.py v4.json out.xlsx ar|en
import json, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
D = json.load(open(sys.argv[1])); OUT = sys.argv[2]; AR = sys.argv[3] == 'ar'
def T(en, ar): return ar if AR else en
NAVY = '1F3864'; GREY = '595959'
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
def F(**k): k.setdefault('name', 'Arial'); k.setdefault('size', 9); return Font(**k)
wb = Workbook(); first = [True]
COL = [('P0', '9C0006', 'FFC7CE'), ('P1', '7F4F00', 'FFEB9C'), ('P2', '1F3864', 'DDEBF7'), ('P3', '375623', 'E2EFDA'), ('NOT VERIFIED', '7F4F00', 'FFEB9C'), ('PASS', '375623', 'E2EFDA'), ('FAIL', '9C0006', 'FFC7CE'),
       ('NOT_PUBLISHED', '9C0006', 'FFC7CE'), ('PARTIALLY_PUBLISHED', '7F4F00', 'FFEB9C'), ('OPEN CONFLICT', '9C0006', 'FFC7CE'), ('ESCALATION', '7F4F00', 'FFEB9C'), ('BLOCKED', '7F4F00', 'FFEB9C'),
       ('VERIFIED', '375623', 'E2EFDA'), ('MATERIAL', '9C0006', 'FFC7CE'), ('FULLY_REVIEWED', '375623', 'E2EFDA'), ('PARTIALLY_REVIEWED', '7F4F00', 'FFEB9C')]
SUB = T('PMU-ADM-360-MASTER-20261022 · v4.0 · 22 Oct 2026 · Prepared by Nader Al Farhan, Administrative Assistant, Student Affairs, Admission Office · Prepared for Management Review',
        'PMU-ADM-360-MASTER-20261022 · الإصدار 4.0 · 22 أكتوبر 2026 · إعداد: نادر آل فرحان، مساعد إداري، شؤون الطلاب، مكتب القبول · معد للمراجعة الإدارية')
def sheet(ten, tar, cols, rows, widths, links=(), center=(), ltr=(), colour=(), note=('', '')):
    ws = wb.active if first[0] else wb.create_sheet(); first[0] = False
    ws.title = T(ten, tar)[:31]; ws.sheet_view.rightToLeft = AR
    ws['A1'] = T(ten, tar); ws['A1'].font = F(size=14, bold=True, color=NAVY)
    ws['A2'] = SUB if not note[0] else T(note[0], note[1]); ws['A2'].font = F(size=8.5, italic=True, color=GREY)
    for j, c in enumerate(cols, 1):
        x = ws.cell(4, j, T(c[0], c[1])); x.font = F(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor=NAVY); x.border = B
        x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center', readingOrder=2 if AR else 1)
    for i, row in enumerate(rows, 5):
        for j, v in enumerate(row, 1):
            x = ws.cell(i, j, v); x.border = B; x.font = F(); l = j in ltr
            x.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if j in center else ('left' if (l or not AR) else 'right'), readingOrder=1 if (l or not AR) else 2)
            if j in links and isinstance(v, str) and v.startswith('http'): x.hyperlink = v; x.font = F(color='0563C1', underline='single')
            if i % 2 == 0: x.fill = PatternFill('solid', fgColor='F7F9FC')
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 36
    if rows: ws.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{4 + len(rows)}'
    ws.freeze_panes = 'B5' if len(cols) > 3 else 'A5'
    for c in colour:
        L = get_column_letter(c)
        for v, fg, bg in COL:
            ws.conditional_formatting.add(f'{L}5:{L}{4 + max(1, len(rows))}', FormulaRule(formula=[f'ISNUMBER(SEARCH("{v}",{L}5))'], font=Font(bold=True, color=fg), fill=PatternFill('solid', fgColor=bg), stopIfTrue=True))
    ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = '4:4'
    return ws
R = D['rows']; S = D['stats']; C = D['coverage22']; DS = D['discovery22']
# 1 Executive Summary
am = collections.Counter(a['code'] for a in D['amend'])
summ = [(T('Title', 'العنوان'), 'PMU Admissions — Management Review & Decision Package · Master Register'), (T('Reference / version', 'المرجع / الإصدار'), 'PMU-ADM-360-MASTER-20261022 · 4.0'),
        (T('Review date', 'تاريخ المراجعة'), D['obs_at']), (T('Prepared by', 'إعداد'), 'Nader Al Farhan | نادر آل فرحان — Administrative Assistant, Student Affairs, Admission Office'),
        (T('Recipient', 'المستلم'), 'Mr. Radhi Alkhaldi — Associate Director of Enrollment'), (T('Root source', 'المصدر الجذري'), 'https://www.pmu.edu.sa/admission/admission'),
        (T('Active findings', 'الملاحظات الفعالة'), S['n']), ('P0 / P1 / P2 / P3', ' / '.join(str(S['plevel'].get(p, 0)) for p in ('P0', 'P1', 'P2', 'P3'))),
        (T('Proven published conflicts', 'تعارضات منشورة مثبتة'), S['conflicts']), (T('Verified strengths', 'نقاط قوة متحقق منها'), len(D['verified_correct'])),
        (T('Verification', 'التحقق'), ' · '.join(f'{k} {v}' for k, v in S['vstatus'].items())), (T('Retired (merged duplicates)', 'مُدمجة (مكررة)'), ', '.join(x['fid'] for x in D['retired']) or '—'),
        (T('Historical amendments', 'التعديلات السابقة'), f"NOT_PUBLISHED {am.get('NOT_PUBLISHED', 0)} · PARTIALLY_PUBLISHED {am.get('PARTIALLY_PUBLISHED', 0)} · PUBLISHED {am.get('PUBLISHED', 0)}"),
        (T('Decisions required / corrective actions', 'القرارات المطلوبة / الإجراءات التصحيحية'), f"{len(D['decisions'])} / {len(D['plan'])}"),
        (T('Full coverage rate', 'نسبة التغطية الكاملة'), f"{C['full_rate']}% ({C['fully']}/{C['eligible']})"),
        (T('Management status', 'الحالة الإدارية'), 'READY FOR MANAGEMENT REVIEW'), (T('Closure status', 'حالة الإغلاق'), 'NOT READY FOR VERIFIED CLOSURE')]
sheet('Executive Summary', 'الملخص التنفيذي', [('Item', 'البند'), ('Value', 'القيمة')], summ, [36, 110], ltr=())
sheet('Source Register', 'سجل المصادر', [('Source ID', 'المعرف'), ('Title', 'العنوان'), ('Origin', 'الأصل'), ('Date / version', 'التاريخ / النسخة'), ('Authority / weight', 'السلطة / الوزن'), ('Completeness / relevance', 'الاكتمال / الصلة'), ('Access status', 'حالة الوصول')],
      [list(s) for s in D['sources']], [9, 58, 26, 22, 34, 46, 18], center=(1, 7))
# 2 Inventory
SN = D['sect']
sheet('Resource Inventory', 'جرد الموارد', [('Resource ID', 'المعرف'), ('Title', 'العنوان'), ('Original URL', 'الرابط الأصلي'), ('Final URL', 'الرابط النهائي'), ('Parent', 'المصدر الأب'), ('Discovery path', 'مسار الاكتشاف'),
      ('Artifact type', 'نوع المورد'), ('Language', 'اللغة'), ('Section', 'القسم'), ('Discovery date', 'تاريخ الاكتشاف'), ('Review date', 'تاريخ المراجعة'), ('Inclusion', 'الإدراج'), ('Review status', 'حالة المراجعة'),
      ('Language pair', 'الزوج اللغوي'), ('Duplicate / canonical', 'التكرار / الأصل'), ('Notes', 'ملاحظات')],
      [[i['rid'], i['title'], i['url'], i['final'], i['parent'], i['found'], i['type'], i['lang'], T(SN[i['section']][0], SN[i['section']][1]), i['disc_date'], i['rev_date'], i['incl'], i['rstatus'],
        i['pair'], i['canon'], T((i['reason_en'] + ('. Not reviewed: ' + i['unchecked_en'] if i['unchecked_en'] else '')), (i['reason_ar'] + ('. غير مفحوص: ' + i['unchecked_ar'] if i['unchecked_ar'] else '')))] for i in D['inventory']],
      [11, 30, 38, 38, 10, 28, 15, 15, 20, 11, 11, 15, 17, 24, 26, 40], links=(3, 4), center=(1, 5, 10, 11, 13), ltr=(2, 3, 4), colour=(13,))
# 3 Findings
rows = []
for r in R:
    rows.append([r['fid'], r['rid'], T(r['page_en'], r['page_ar']), r['url'] if r['url'].startswith('http') else '', r['artifact'], r['lang_res'], T(r.get('location_en', ''), r['location']), r['current_text'],
                 T(r['state22_en'], r['state22_ar']), T(r['cls_en'], r['cls_ar']), r['verified_on'], ', '.join(r['eids']), T(r['auth_ref'], r['auth_ref_ar']), T(r['authority_holder'], r['authority_holder_ar']),
                 T(r['impact_en'], r['impact_ar']), r['plevel'], T(r['rationale_en'], r['rationale_ar']), T(r['description_en'], r['description']),
                 T(r['recommendation_en'], r.get('recommendation_ar') or r['recommendation']), ', '.join(r['dids']), r['owner_code'], r['vstatus'], r['res_status'], r['pub_status'],
                 (r.get('followup_en') or '') + ((' | ' + r['challenge_note']) if r.get('challenge_note') else ''), r['closure_ev']])
sheet('Findings', 'الملاحظات', [('Finding ID', 'المعرف'), ('Resource ID', 'معرف المورد'), ('Page title', 'عنوان الصفحة'), ('URL', 'الرابط'), ('Artifact type', 'نوع المورد'), ('Language', 'اللغة'), ('Exact location', 'الموضع'),
      ('Observed claim (verbatim)', 'النص المرصود (حرفيًا)'), ('Current state', 'الحالة الحالية'), ('Classification', 'التصنيف'), ('Date verified', 'تاريخ التحقق'), ('Evidence IDs', 'معرفات الأدلة'),
      ('Authoritative reference', 'المرجع الحاكم'), ('Authority', 'صاحب الصلاحية'), ('Impact', 'الأثر'), ('Severity', 'الخطورة'), ('Severity rationale', 'مبرر الخطورة'), ('Finding', 'الملاحظة'),
      ('Recommendation', 'التوصية'), ('Decision required', 'القرار المطلوب'), ('Owner', 'المسؤول'), ('Verification', 'التحقق'), ('Resolution status', 'حالة الحسم'), ('Publication status', 'حالة النشر'),
      ('Follow-up', 'المتابعة'), ('Closure evidence', 'دليل الإغلاق')],
      rows, [8, 10, 22, 28, 12, 8, 22, 44, 22, 22, 11, 14, 34, 24, 16, 7, 30, 50, 44, 9, 22, 11, 20, 24, 40, 26], links=(4,), center=(1, 2, 6, 11, 16, 20, 22), ltr=(4, 8, 12), colour=(16, 22, 23))
# 4 Evidence
ev = [[x['eid'], x['fid'], x['url'], x['excerpt'], T(x['locator'], x['locator_ar']), x['first_observed'], x['observed'], x['pub_date'], x['shot'], x['page_shot']] for x in D['evidence']]
ev += [[f"RC-{x['id']}", x['id'].rstrip('bc'), '', x['snippet'], T('Live re-confirmation', 'إعادة تأكيد حية'), '', '2026-10-22', f"HTTP {x['http']} · matches {x['hits']}", '', ''] for x in D['recheck']]
sheet('Evidence', 'الأدلة', [('Evidence ID', 'معرف الدليل'), ('Finding ID', 'الملاحظة'), ('Source URL', 'رابط المصدر'), ('Excerpt (verbatim)', 'المقتطف (حرفيًا)'), ('Locator / context', 'الموضع / السياق'),
      ('First observed', 'أول رصد'), ('Last observed', 'آخر رصد'), ('Source date / result', 'تاريخ المصدر / النتيجة'), ('Annotated screenshot', 'لقطة مظللة'), ('Original capture', 'اللقطة الأصلية')],
      ev, [12, 9, 40, 56, 28, 12, 11, 24, 18, 18], links=(3,), center=(1, 2, 6, 7), ltr=(3, 4, 9, 10))
# 5 Coverage
cov = [(T('Discovered link instances', 'روابط مرصودة'), C['discovered_link_instances']), (T('Discovered unique URLs', 'روابط فريدة مكتشفة'), C['discovered_unique_urls']),
       (T('Root-page placements', 'مواضع الصفحة الجذرية'), C['root_placements']), (T('Unique destinations', 'الوجهات الفريدة'), C['unique_destinations']), (T('Eligible resources', 'الموارد المؤهلة'), C['eligible']),
       ('FULLY_REVIEWED', C['fully']), ('PARTIALLY_REVIEWED', C['partially']), ('BLOCKED', C['blocked']), ('NOT_REVIEWED', C['not_reviewed']), ('EXCLUDED', C['excluded']),
       (T('PDFs / PDF pages', 'ملفات PDF / صفحاتها'), f"{C['pdfs']} / {C['pdf_pages']}"), (T('Arabic standalone resources', 'موارد عربية مستقلة'), C['arabic_standalone']),
       (T('English resources', 'موارد إنجليزية'), C['english']), (T('Confirmed AR/EN pairs (bilingual resources / segment pairs)', 'أزواج عربية/إنجليزية مؤكدة (موارد / مقاطع)'), f"{C['pairs']} / {C['segment_pairs']}"),
       (T('Full Coverage Rate = FULLY_REVIEWED / ELIGIBLE', 'نسبة التغطية الكاملة = FULLY_REVIEWED / المؤهلة'), f"{C['full_rate']}%"),
       (T('Link status (unique URLs)', 'حالة الروابط (فريدة)'), ' · '.join(f'{k}: {v}' for k, v in sorted(DS['link_codes'].items()))), (T('HTTP 404 links', 'روابط 404'), ' | '.join(DS['broken']))]
sheet('Coverage', 'التغطية', [('Measure', 'المقياس'), ('Value', 'القيمة')], cov, [60, 90], note=('Discovery and capture of 22 Oct 2026; no number is inherited from earlier rounds. Status 000 = no response from the review network (external/internal hosts), not counted as broken.', 'اكتشاف والتقاط 22 أكتوبر 2026؛ لا رقم موروث من جولات سابقة. الحالة 000 = لا استجابة من شبكة المراجعة (نطاقات خارجية/داخلية) ولا تُحسب معطلة.'))
# 6 Language parity
sheet('Language Parity', 'التكافؤ اللغوي', [('Pair', 'الزوج'), ('AR resource', 'المورد العربي'), ('EN resource', 'المورد الإنجليزي'), ('Topic', 'الموضوع'), ('English (verbatim)', 'الإنجليزية (حرفيًا)'), ('Arabic (verbatim)', 'العربية (حرفيًا)'),
      ('Numeric difference', 'فرق رقمي'), ('Policy difference', 'فرق في السياسة'), ('Semantic impact', 'الأثر الدلالي'), ('Finding', 'الملاحظة'), ('Evidence', 'الدليل'), ('Status', 'الحالة')],
      [[p['pid'], p['ar_res'], p['en_res'], p['topic'], p['en'], p['ar'], p['numeric'], p['policy'], p['semantic'], p['fid'], p['evidence'], p['status']] for p in D['parity']] +
      [['L-NOTE', '', '', T('Separate Arabic version', 'نسخة عربية مستقلة'), '', '', '', '', T(D['lang_note'][0], D['lang_note'][1]), '', '', 'NOT FOUND WITHIN REVIEWED INVENTORY']],
      [8, 22, 22, 24, 40, 40, 12, 30, 34, 9, 22, 22], center=(1, 10), ltr=(5,), colour=(12,))
# 7 Strengths
sheet('Strengths', 'نقاط القوة', [('ID', 'المعرف'), ('Page', 'الصفحة'), ('URL', 'الرابط'), ('Evidence (verbatim)', 'الدليل (حرفيًا)'), ('Strength', 'نقطة القوة'), ('Classification', 'التصنيف'), ('Verified', 'التحقق')],
      [[v['vid'], T(v['page_en'], v['page_ar']), v['url'], v['excerpt'], T(v['en'], v['ar']), T(*D['c12']['VERIFIED_STRENGTH']), 'VERIFIED' if v['verified'] else 'NOT VERIFIED'] for v in D['verified_correct']],
      [8, 24, 34, 46, 60, 18, 12], links=(3,), center=(1, 7), ltr=(3, 4), colour=(7,))
# 8 Historical amendments
sheet('Historical Amendments', 'التعديلات السابقة', [('ID', 'المعرف'), ('Requested amendment (26 Oct 2025)', 'التعديل المطلوب (26 أكتوبر 2025)'), ('Classification', 'التصنيف'), ('Status', 'الحالة'), ('Current published evidence', 'الدليل المنشور الحالي'), ('Decision', 'القرار'), ('Checked', 'تاريخ التحقق'), ('Basis', 'الأساس')],
      [[a['hid'], T(a['req_en'], a['req_ar']), T(*D['c12']['HISTORICAL_CHANGE_NOT_PUBLISHED']), a['code'], T(a['ev_en'], a['ev_ar']), a['did'], a['checked'], T(a['basis_en'], a['basis_ar'])] for a in D['amend']],
      [7, 44, 22, 22, 66, 9, 12, 40], center=(1, 4, 6, 7), colour=(4,))
# 9 Decisions
sheet('Decisions', 'القرارات', [('Decision ID', 'المعرف'), ('Finding IDs', 'الملاحظات'), ('Question', 'السؤال'), ('Evidence', 'الأدلة'), ('Impact', 'الأثر'), ('Risk if deferred', 'المخاطر عند التأجيل'), ('Options', 'الخيارات'),
      ('Recommendation', 'التوصية'), ('Reference requiring approval', 'المرجع المطلوب اعتماده'), ('Required authority', 'صاحب الصلاحية'), ("Mr. Radhi's role", 'دور الأستاذ راضي'), ('Implementation owner', 'مسؤول التنفيذ'),
      ('Publication owner', 'مسؤول النشر'), ('Verification owner', 'مسؤول التحقق'), ('Status', 'الحالة'), ('Management direction', 'توجيه الإدارة'), ('Date', 'التاريخ')],
      [[d['did'], ', '.join(d['fids']) if d['did'] != 'D-09' else T(f"{len(d['fids'])} editorial findings + {len(D['edits'])} corrections", f"{len(d['fids'])} ملاحظة تحريرية + {len(D['edits'])} تصحيحًا"), T(d['t_en'] + ' — ' + d['q_en'], d['t_ar'] + ' — ' + d['q_ar']),
        d['evidence'], d['impact_en'], T(d['risk_en'], d['risk_ar']), T(d['alt_en'], d['alt_ar']), T(d['rec_en'], d['rec_ar']), T(d['ref_en'], d['ref_ar']), T(d['auth'], d['auth_ar']), T(d['role'], d['role_ar']),
        T(d['impl'], d['impl_ar']), T(d['publisher'], d['publisher_ar']), T(d['ver'], d['ver_ar']), T('PENDING — NOT APPROVED', 'معلّق — غير معتمد'), '', ''] for d in D['decisions']],
      [8, 26, 50, 22, 16, 34, 36, 40, 32, 26, 20, 24, 24, 20, 18, 22, 11], center=(1, 15), ltr=(2, 4))
# 10 Corrective actions
sheet('Corrective Actions', 'الإجراءات التصحيحية', [('Action ID', 'المعرف'), ('Finding IDs', 'الملاحظات'), ('Decision ID', 'القرار'), ('Severity', 'الخطورة'), ('Current state', 'الحالة الحالية'), ('Target state', 'الحالة المستهدفة'),
      ('Required change', 'التغيير المطلوب'), ('Governing reference', 'المرجع الحاكم'), ('Approval required', 'الاعتماد المطلوب'), ('Owner', 'المسؤول'), ('Dependency', 'الاعتمادية'), ('Target date (proposal)', 'الموعد (مقترح)'),
      ('Status', 'الحالة'), ('Publication evidence', 'دليل النشر'), ('Re-verification', 'إعادة التحقق'), ('Closure status', 'حالة الإغلاق')],
      [[p['aid'], p['fids'], p['did'], p['prio'], p['problem'], p['outcome'], T(p['en'], p['ar']), T(p['gov_ref'], p['gov_ref_ar']), p['approval'], p['owner'], p['dependency'], p['target'], p['status'], '', '', p['closure']] for p in D['plan']],
      [8, 20, 9, 8, 36, 30, 40, 34, 28, 26, 26, 26, 14, 18, 16, 12], center=(1, 3, 4, 13, 16), ltr=(2,), colour=(4,),
      note=('Lifecycle after authorisation: APPROVE → IMPLEMENT → SAVE → REOPEN → READ BACK → COMPARE → VERIFY → CLOSE. Target dates are proposals and commit no unit.', 'دورة ما بعد التفويض: اعتماد ← تنفيذ ← حفظ ← إعادة فتح ← إعادة قراءة ← مقارنة ← تحقق ← إغلاق. المواعيد مقترحة ولا تُلزم أي جهة.'))
# 11 Follow-up
sheet('Follow-up', 'المتابعة', [('ID', 'المعرف'), ('Follow-up step', 'خطوة المتابعة'), ('Responsible', 'المسؤول'), ('Status', 'الحالة'), ('Date', 'التاريخ'), ('Evidence', 'الدليل')],
      [[f[0], T(f[1], f[2]), f[3], f[4], '', ''] for f in D['followup']], [8, 70, 26, 18, 12, 30], center=(1, 4, 5))
sheet('Applicant Journey', 'رحلة المتقدم', [('ID', 'المعرف'), ('Element', 'العنصر'), ('Findable', 'متاح'), ('Observation', 'الملاحظة'), ('Findings', 'الملاحظات')],
      [[j['jid'], T(j['en'], j['ar']), T(j['f_en'], j['f_ar']), T(j['o_en'], j['o_ar']), j['fids']] for j in D['journey']], [8, 24, 12, 70, 30], center=(1, 3), ltr=(5,),
      note=('Published applicant information journey — findability only; no applicant data or opinion.', 'رحلة المتقدم في المعلومات المنشورة — قابلية العثور فقط؛ دون بيانات أو آراء متقدمين.'))
# 12 QC
qrows = [[g['gid'], T(g['en'], g['ar']), g['result'], g['evidence'], g.get('residual', ''), g.get('action', '')] for g in D.get('gates4', [])]
qrows += [[f"IC-{c['fid']}", T('Independent challenge', 'تحدٍّ مستقل'), c['verdict'], c.get('rationale', ''), D['challenge_res'].get(c['fid'], ''), ''] for c in D.get('challenge', [])]
sheet('QC', 'ضبط الجودة', [('ID', 'المعرف'), ('Gate / check', 'المعيار / الفحص'), ('Result', 'النتيجة'), ('Evidence', 'الدليل'), ('Residual / resolution', 'المتبقي / المعالجة'), ('Action', 'الإجراء')],
      qrows, [12, 40, 18, 70, 60, 30], center=(1, 3), colour=(3,))
sheet('Editorial Corrections', 'التصحيحات التحريرية', [('Change ID', 'المعرف'), ('URL', 'الرابط'), ('Current text', 'النص الحالي'), ('Proposed text', 'النص المقترح'), ('Verification', 'التحقق')],
      [[x['cid'], x['url'], x['before'], x['after'], 'VERIFIED' if x['verified'] else 'BLOCKED (dynamic portal label)'] for x in D['edits']], [10, 44, 40, 40, 26], links=(2,), center=(1, 5), ltr=(2, 3, 4), colour=(5,),
      note=('Ready editorial corrections (D-09); none changes a value, requirement or fee.', 'تصحيحات تحريرية جاهزة (D-09)؛ لا يغيّر أيٌّ منها قيمة أو شرطًا أو رسمًا.'))
# 13 Definitions
defs = [(k, T(*v)) for k, v in D['c12'].items()] + [(k, T(*v)) for k, v in D['pl'].items()] + [
 ('FULLY_REVIEWED', T('All eligible content of the resource reviewed.', 'رُوجع كل المحتوى المؤهل في المورد.')), ('PARTIALLY_REVIEWED', T('A defined part reviewed; remainder recorded.', 'رُوجع جزء محدد مع تسجيل الباقي.')),
 ('BLOCKED', T('Could not be completed with permitted means; reason recorded.', 'تعذر الإكمال بالوسائل المسموح بها؛ والسبب مسجل.')), ('NOT_REVIEWED', T('Eligible but not yet reviewed.', 'مؤهل ولم يُراجع بعد.')),
 ('EXCLUDED', T('Outside scope with a documented reason.', 'خارج النطاق بسبب موثق.')), ('VERIFIED', T('Observed state confirmed against retrievable evidence on the stated date.', 'تأكدت الحالة المرصودة بدليل قابل للاسترجاع في التاريخ المذكور.')),
 ('OPEN CONFLICT', T('Two published statements conflict; no governing reference resolves them.', 'بيانان منشوران متعارضان دون مرجع حاكم يحسم.')), ('ESCALATION REQUIRED', T('High-impact matter needing a decision by a competent authority.', 'مسألة عالية الأثر تحتاج قرار صاحب صلاحية.')),
 ('OWNER_TO_CONFIRM', T('No evidence identifies the responsible unit.', 'لا دليل يحدد الجهة المسؤولة.')), ('DIRECT / REFER', T('Within the Admissions Office remit / to be referred to the competent authority.', 'ضمن صلاحية مكتب القبول / يُحال إلى صاحب الصلاحية.')),
 ('EDITORIAL DIFFERENCE / MATERIAL PARITY ISSUE', T('Arabic–English difference without / with effect on meaning or policy.', 'اختلاف عربي–إنجليزي دون / مع أثر في المعنى أو السياسة.')),
 ('READY FOR MANAGEMENT REVIEW', T('Package fit for management decision; gaps disclosed.', 'الحزمة صالحة لقرار الإدارة مع إعلان الفجوات.')), ('NOT READY FOR VERIFIED CLOSURE', T('A decision, reference, conflict, gap or unverified correction remains.', 'يبقى قرار أو مرجع أو تعارض أو فجوة أو تصحيح غير متحقق.')),
 ('VERIFIED-OUTCOME-CLOSED', T('Only after DECISION → APPROVAL → IMPLEMENTATION → PUBLICATION → READ-BACK → VERIFY → REGRESSION → ACCEPTANCE.', 'فقط بعد: قرار ← اعتماد ← تنفيذ ← نشر ← إعادة قراءة ← تحقق ← فحص ارتدادي ← قبول.'))]
sheet('Definitions', 'التعريفات', [('Term', 'المصطلح'), ('Definition', 'التعريف')], defs, [40, 110])
# 14 Change log (+ retired)
sheet('Change Log', 'سجل التغييرات', [('Version', 'الإصدار'), ('Date', 'التاريخ'), ('Change', 'التغيير')],
      [list(c) for c in D['changelog']] + [[f"retired {x['fid']}", '2026-10-22', x['reason']] for x in D['retired']], [12, 12, 130], center=(1, 2))
wb.properties.creator = 'Nader Al Farhan'; wb.properties.lastModifiedBy = 'Nader Al Farhan'; wb.properties.title = 'PMU Admissions — Master Register'; wb.properties.subject = 'Admissions Office content review'
wb.properties.description = None; wb.properties.keywords = None
wb.save(OUT)
import zipfile, shutil, os, re as _re
tmp = OUT + '.tmp'
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == 'docProps/app.xml':
            data = _re.sub(rb'<Application>.*?</Application>', b'<Application>Microsoft Excel</Application>', data)
            data = _re.sub(rb'<AppVersion>.*?</AppVersion>', b'<AppVersion>16.0300</AppVersion>', data)
        zo.writestr(it, data)
os.replace(tmp, OUT); print('saved', OUT, len(wb.sheetnames))
