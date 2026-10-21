# Build the bilingual 360 register workbook. Usage: b360_xlsx.py data.json out.xlsx en|ar
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
D = json.load(open(sys.argv[1])); OUT = sys.argv[2]; L = sys.argv[3]
AR = L == 'ar'
def T(en, ar): return ar if AR else en
FONT = 'Arial'; NAVY = '1F3864'; GREY = '595959'; LIGHT = 'EAF1FB'
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
def F(**k): k.setdefault('name', FONT); k.setdefault('size', 10); return Font(**k)
wb = Workbook(); first = True
def sheet(title_en, title_ar, heading_en, heading_ar, sub_en, sub_ar, cols, rows, widths, link_cols=(), center=(), ltr_cols=(), prio_col=None):
    global first
    ws = wb.active if first else wb.create_sheet(); first = False
    ws.title = T(title_en, title_ar)[:31]; ws.sheet_view.rightToLeft = AR
    ws['A1'] = T(heading_en, heading_ar); ws['A1'].font = F(size=14, bold=True, color=NAVY)
    ws['A2'] = T(sub_en, sub_ar); ws['A2'].font = F(size=9, italic=True, color=GREY)
    hdr = [T(c[0], c[1]) for c in cols]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center', readingOrder=2 if AR else 1); c.border = B
    for i, row in enumerate(rows, 5):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.border = B; c.font = F(size=9)
            ltr = j in ltr_cols
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if j in center else ('left' if (ltr or not AR) else 'right'), readingOrder=1 if ltr or not AR else 2)
            if j in link_cols and isinstance(v, str) and v.startswith('http'):
                c.hyperlink = v; c.font = F(size=9, color='0563C1', underline='single')
        if i % 2 == 0:
            for j in range(1, len(row) + 1):
                if ws.cell(i, j).fill.fgColor.rgb in (None, '00000000'): ws.cell(i, j).fill = PatternFill('solid', fgColor='F7F9FC')
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 32
    if rows:
        ws.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{4 + len(rows)}'
    ws.freeze_panes = 'B5' if len(cols) > 3 else 'A5'
    if prio_col:
        col = get_column_letter(prio_col)
        for v, fg, bg in [('P0', '9C0006', 'FFC7CE'), ('P1', '7F4F00', 'FFEB9C'), ('P2', '1F3864', 'DDEBF7'), ('P3', '375623', 'E2EFDA')]:
            ws.conditional_formatting.add(f'{col}5:{col}{4 + max(1, len(rows))}', FormulaRule(formula=[f'{col}5="{v}"'], font=Font(bold=True, color=fg), fill=PatternFill('solid', fgColor=bg)))
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '4:4'
    return ws

R = D['rows']; S = D['stats']; C = D['coverage']
# 1 Summary
summ = [
 (T('Observation date and time', 'تاريخ الرصد ووقته'), D['obs_at']),
 (T('Root URL', 'الرابط الجذري'), 'https://www.pmu.edu.sa/admission/admission'),
 (T('Root access status', 'حالة الوصول إلى الجذر'), 'HTTP 200'),
 (T('Scope', 'النطاق'), T('Admissions Office page, its five sections (Future Students; International Students’ Office; Graduate Admissions; Undergraduate Admissions; Undergraduate Placement Tests), their direct destinations and the files/fee pages needed to verify them.', 'صفحة مكتب القبول وأقسامها الخمسة (الطلاب المستقبليون؛ مكتب الطلاب الدوليين؛ القبول في الدراسات العليا؛ القبول في البكالوريوس؛ اختبارات تحديد المستوى) ووجهاتها المباشرة والملفات وصفحات الرسوم اللازمة للتحقق.')),
 (T('Total findings (open / appear resolved)', 'إجمالي الملاحظات (مفتوحة / يبدو أنها عولجت)'), f"{S['n']} ({S['open']} / {S['n'] - S['open']})"),
 ('P0 / P1 / P2 / P3', f"{S['plevel'].get('P0',0)} / {S['plevel'].get('P1',0)} / {S['plevel'].get('P2',0)} / {S['plevel'].get('P3',0)}"),
 (T('Proven published conflicts', 'تعارضات منشورة مثبتة'), S['conflicts']),
 (T('Proven strengths', 'نقاط قوة مثبتة'), len(D['strengths'])),
 (T('Institutional decisions requested', 'قرارات مؤسسية مطلوبة'), len(D['decisions'])),
 (T('Ready editorial corrections (annex)', 'تصحيحات تحريرية جاهزة (مرفق)'), f"{len(D['edits'])} ({sum(e['verified'] for e in D['edits'])} {T('verified against published text', 'مُتحقق منها في النص المنشور')})"),
 (T('Root-page link placements / unique destinations', 'روابط الصفحة الجذرية / الوجهات الفريدة'), f"{C['root_placements']} / {C['root_unique']}"),
 (T('Eligible resources: fully / partially examined', 'الموارد المؤهلة: مفحوصة بالكامل / جزئيًا'), f"{C['eligible']}: {C['full']} / {C['partial']}"),
 (T('Coverage of eligible resources (full | any)', 'تغطية الموارد المؤهلة (كاملة | أي فحص)'), f"{C['pct_full']}% | {C['pct_any']}%"),
 (T('Link instances / unique URLs checked', 'روابط مرصودة / روابط فريدة فُحصت'), f"{C['links_instances']} / {C['links_unique']}"),
 (T('PDF files / pages read', 'ملفات PDF / صفحاتها المقروءة'), f"{C['pdfs']} / {C['pdf_pages']}"),
 (T('Language pairs compared', 'الأزواج اللغوية المقارنة'), T('1 bilingual page (/apply, 9 segment pairs) + 1 bilingual form (1 pair); no separate Arabic site version found', 'صفحة ثنائية اللغة واحدة (/apply، 9 أزواج مقاطع) + نموذج ثنائي اللغة (زوج واحد)؛ ولم يُعثر على نسخة عربية مستقلة')),
 (T('Overall audit status', 'الحالة العامة للتدقيق'), 'NOT READY FOR VERIFIED CLOSURE'),
 (T('Reason for status', 'سبب الحالة'), T('Policy values in conflict await institutional confirmation; portal registration labels and one link could not be verified; no applicant data or feedback was available; no corrections have been published.', 'القيم المتعارضة تنتظر تأكيدًا مؤسسيًا؛ وتعذر التحقق من تسميات تسجيل البوابة ورابط واحد؛ ولا تتوفر بيانات أو آراء للمتقدمين؛ ولم يُنشر أي تصحيح.')),
 (T('Package readiness for management review', 'جاهزية الحزمة للعرض الإداري'), T('Ready for management review (draft, not approved)', 'جاهزة للعرض الإداري (مسودة غير معتمدة)')),
 (T('Implementation of corrections', 'تنفيذ التصحيحات'), T('Outside this mandate; nothing changed on the website', 'خارج صلاحية هذا التكليف؛ لم يُعدَّل شيء في الموقع')),
]
sheet('Summary', 'الملخص', 'PMU Admissions Office — Content Review Register', 'سجل مراجعة محتوى مكتب القبول — جامعة الأمير محمد بن فهد',
      'Prepared for management review. Not an approval, publication confirmation or closure certificate.', 'معد للمراجعة الإدارية. لا يمثل اعتمادًا أو تأكيد نشر أو شهادة إغلاق.',
      [('Indicator', 'المؤشر'), ('Value', 'القيمة')], summ, [46, 110])

# Definitions
defs = [('FULL', T('Fully examined: all eligible content, including informative visual elements, reviewed.', 'مفحوص بالكامل: فُحص المحتوى المؤهل كله بما فيه العناصر المرئية المؤثرة.')),
        ('PARTIAL', T('Partially examined: a defined part reviewed; remaining parts recorded.', 'مفحوص جزئيًا: فُحص جزء محدد مع تسجيل الأجزاء المتبقية.')),
        ('UNREACHABLE', T('Could not be retrieved after appropriate, different attempts.', 'تعذر الوصول بعد محاولات مناسبة ومختلفة.')),
        ('NOT EXAMINED', T('Discovered and eligible, examination not started.', 'مورد مكتشف ومؤهل لم يبدأ فحصه.')),
        ('EXCLUDED', T('Outside scope, with a documented reason.', 'خارج النطاق مع سبب موثق.'))] + \
       [(k, T(v[0], v[1])) for k, v in D['pl'].items()] + [(k, T(v[0], v[1])) for k, v in D['cat'].items()] + \
       [('Re-verification', T('Each finding’s verbatim text was re-matched against the 21 Oct 2026 capture; changed items are flagged.', 'طوبق النص الحرفي لكل ملاحظة مع لقطة 21 أكتوبر 2026؛ ووُسمت البنود المتغيرة.'))]
sheet('Definitions', 'التعريفات', 'Status, priority and category definitions', 'تعريفات الحالات والأولويات والتصنيفات', '', '', [('Term', 'المصطلح'), ('Definition', 'التعريف')], defs, [22, 120])

# 2 Inventory
SN = D['sect']
inv_rows = [[i['rid'], i['title'], i['type'], i['lang'], i['url'], i['final'], i['found'], T(SN[i['section']][0], SN[i['section']][1]), i['incl'], T(i['reason_en'], i['reason_ar']),
             i['at'], i['http'], i['status'], T(i['unchecked_en'], i['unchecked_ar']), i['rel']] for i in D['inventory']]
sheet('Resource Inventory', 'جرد الموارد', 'Resource inventory', 'جرد الموارد', 'Stable IDs; Arabic and English are never merged as one resource.', 'معرفات ثابتة؛ ولا تُدمج العربية والإنجليزية في مورد واحد.',
      [('ID', 'المعرف'), ('Title', 'العنوان'), ('Type', 'النوع'), ('Language', 'اللغة'), ('Original URL', 'الرابط الأصلي'), ('Final URL', 'الرابط النهائي'), ('Discovered from', 'مسار الاكتشاف'),
       ('Section', 'القسم'), ('Inclusion', 'الإدراج'), ('Reason', 'السبب'), ('Accessed', 'تاريخ الوصول'), ('HTTP', 'HTTP'), ('Exam status', 'حالة الفحص'), ('Parts not examined', 'الأجزاء غير المفحوصة'), ('Relation', 'العلاقة')],
      inv_rows, [11, 34, 16, 16, 40, 40, 30, 22, 16, 36, 11, 7, 11, 36, 18], link_cols=(5, 6), center=(1, 11, 12, 13), ltr_cols=(2, 5, 6))

# 3 Findings
f_rows = []
for r in R:
    f_rows.append([r['fid'], T(r['title_en'], r['title_ar']), T(r['page_en'], r['page_ar']), r['url'] if r['url'].startswith('http') else '', T(r.get('location_en', ''), r['location']),
                   r['current_text'], T(r['cat_en'], r['cat_ar']), r['plevel'], T(r['rationale_en'], r['rationale_ar']), T(r['description_en'], r['description']),
                   T(r.get('conflict_statement', ''), r.get('conflict_ar', '')), T(r['recommendation_en'], r.get('recommendation_ar') or r['recommendation']),
                   ', '.join(r['dids']), T(r['owner_en'], r['owner_ar']), ', '.join(r['eids']), T(r['rv_en'], r['rv_ar']), T(r['rv_note_en'], r['rv_note_ar']),
                   T('Open — awaiting management direction', 'مفتوحة — بانتظار توجيه الإدارة') if r['rv'] != 'RESOLVED' else T('Appears resolved — owner to confirm', 'يبدو أنها عولجت — يؤكدها المالك')])
sheet('Findings', 'الملاحظات', 'Findings register', 'سجل الملاحظات', 'Published text is quoted verbatim in its original language. Recommendations are proposals, not approved decisions.',
      'يُقتبس النص المنشور حرفيًا بلغته الأصلية. التوصيات مقترحات وليست قرارات معتمدة.',
      [('ID', 'المعرف'), ('Title', 'العنوان'), ('Page', 'الصفحة'), ('URL', 'الرابط'), ('Location', 'الموقع'), ('Published text (verbatim)', 'النص المنشور (حرفيًا)'), ('Category', 'التصنيف'),
       ('Priority', 'الأولوية'), ('Priority rationale', 'مبرر الأولوية'), ('Finding and impact', 'الملاحظة وأثرها'), ('Conflict statement', 'بيان التعارض'), ('Proposed action', 'الإجراء المقترح'),
       ('Decision', 'القرار'), ('Proposed owner', 'الجهة المقترحة'), ('Evidence IDs', 'معرفات الأدلة'), ('Verification (21 Oct 2026)', 'التحقق (21 أكتوبر 2026)'), ('Verification note', 'ملاحظة التحقق'), ('Follow-up status', 'حالة المتابعة')],
      f_rows, [8, 30, 22, 30, 24, 44, 18, 8, 30, 50, 24, 46, 9, 22, 14, 22, 30, 20], link_cols=(4,), center=(1, 8, 13), ltr_cols=(4, 6, 15), prio_col=8)

# 4 Evidence
e_rows = [[e['eid'], e['fid'], e['url'], e['excerpt'], T(e['locator'], e['locator_ar']), e['first_observed'], e['observed'], e['pub_date'], e['shot'], e['page_shot']] for e in D['evidence']]
sheet('Evidence', 'الأدلة', 'Evidence register', 'سجل الأدلة', 'Observation date is separate from any publication date shown by the source. Screenshots are original captures of the live pages.',
      'تاريخ الرصد منفصل عن تاريخ نشر المصدر. اللقطات أصلية من الصفحات الحية.',
      [('Evidence ID', 'معرف الدليل'), ('Finding', 'الملاحظة'), ('Source URL', 'رابط المصدر'), ('Excerpt (verbatim)', 'المقتطف (حرفيًا)'), ('Locator', 'الموضع'), ('First observed', 'أول رصد'),
       ('Last observed', 'آخر رصد'), ('Source publication/update date', 'تاريخ نشر المصدر/تحديثه'), ('Highlighted screenshot', 'لقطة مظللة'), ('Full-page capture', 'لقطة الصفحة كاملة')],
      e_rows, [11, 8, 40, 60, 30, 12, 12, 24, 18, 20], link_cols=(3,), center=(1, 2, 6, 7), ltr_cols=(3, 4, 9, 10))

# 5 Coverage
cov = [
 (T('Root-page link placements (five sections)', 'روابط الصفحة الجذرية (الأقسام الخمسة)'), C['root_placements'], T('Apply Now and Apply both lead to /apply', '«Apply Now» و«Apply» يقودان إلى /apply')),
 (T('Unique destinations from the root', 'الوجهات الفريدة من الجذر'), C['root_unique'], ''),
 (T('Unique resources in inventory (incl. excluded groups)', 'الموارد الفريدة في الجرد (مع مجموعات المستبعد)'), C['resources_unique'], ''),
 (T('Eligible resources', 'الموارد المؤهلة'), C['eligible'], T('Root + 35 destinations + 9 verification sources + 2 PDFs', 'الجذر + 35 وجهة + 9 مصادر تحقق + ملفا PDF')),
 (T('Fully examined', 'مفحوص بالكامل'), C['full'], ''),
 (T('Partially examined', 'مفحوص جزئيًا'), C['partial'], T('Staff list (personal data), portal (public page only), Aptis guide images, preparatory study plan', 'قائمة الموظفين (بيانات شخصية)، البوابة (الصفحة العامة فقط)، صور دليل Aptis، الخطة الدراسية للتحضيري')),
 (T('Unreachable', 'تعذر الوصول'), C['unreachable'], ''), (T('Not examined', 'لم يُفحص'), C['not_examined'], ''),
 (T('Excluded groups', 'مجموعات مستبعدة'), C['excluded'], T('P01/P02 general pages + 5 excluded link groups', 'الصفحتان العامتان + 5 مجموعات روابط مستبعدة')),
 (T('Coverage — fully examined / eligible', 'التغطية — المفحوص بالكامل / المؤهل'), f"{C['pct_full']}%", T('Denominator includes partial items', 'المقام يشمل البنود الجزئية')),
 (T('Coverage — any examination / eligible', 'التغطية — أي فحص / المؤهل'), f"{C['pct_any']}%", T('Applies to the discovered inventory only; not a claim about the whole website', 'تخص الجرد المكتشف فقط؛ وليست ادعاءً بشأن الموقع كله')),
 (T('Link instances on in-scope pages', 'روابط مرصودة في صفحات النطاق'), C['links_instances'], ''),
 (T('Unique URLs status-checked', 'روابط فريدة فُحصت حالتها'), C['links_unique'], ', '.join(f'{k}: {v}' for k, v in sorted(C['link_codes'].items()))),
 (T('Confirmed HTTP 404 links', 'روابط 404 مؤكدة'), len(C['broken']), ' | '.join(C['broken'])),
 (T('Status 000 (no response from review environment)', 'الحالة 000 (لا استجابة من بيئة المراجعة)'), C['link_codes'].get('000', 0), T('External/social/internal hosts blocked by the review network or malformed hrefs; not treated as public failures', 'نطاقات خارجية/اجتماعية/داخلية محجوبة في شبكة المراجعة أو روابط مشوهة؛ لا تُعد أعطالًا لدى الجمهور')),
 (T('PDF files / pages read', 'ملفات PDF / الصفحات المقروءة'), f"{C['pdfs']} / {C['pdf_pages']}", T('Web Admission Guide 21 pp. (image of p. 20 checked); Aptis Candidate Guide 50 pp. (text)', 'دليل التقديم 21 صفحة (فُحصت صورة الصفحة 20)؛ دليل Aptis 50 صفحة (نص)')),
 (T('Language pairs compared', 'الأزواج اللغوية المقارنة'), 10, T('9 segment pairs on /apply + 1 on the Admission Query form', '9 أزواج مقاطع في /apply + زوج في نموذج الاستفسار')),
]
sheet('Coverage', 'التغطية', 'Coverage and gaps', 'التغطية والفجوات', '', '', [('Measure', 'المقياس'), ('Value', 'القيمة'), ('Note', 'ملاحظة')], cov, [52, 16, 90], center=(2,))

# 6 Language comparison
lp = [[l[0], l[1], 'https://www.pmu.edu.sa/apply' if l[0] != 'L-10' else 'https://pmu.edu.sa/admission/admissions_queries', l[2], l[3], T(l[4], l[5]), l[6]] for l in D['lang_pairs']]
lp.append(['L-NOTE', T('Separate Arabic version', 'النسخة العربية المستقلة'), '', '', '', T(D['lang_note'][0], D['lang_note'][1]), ''])
sheet('Language Comparison', 'المقارنة اللغوية', 'Arabic–English comparison', 'المقارنة بين العربية والإنجليزية', 'Editorial differences are separated from differences that change meaning.', 'يُفصل الاختلاف التحريري عن الاختلاف الذي يغير المعنى.',
      [('Pair ID', 'المعرف'), ('Topic', 'الموضوع'), ('Source', 'المصدر'), ('English (verbatim)', 'الإنجليزية (حرفيًا)'), ('Arabic (verbatim)', 'العربية (حرفيًا)'), ('Result', 'النتيجة'), ('Linked finding', 'الملاحظة المرتبطة')],
      lp, [9, 26, 30, 46, 46, 50, 12], link_cols=(3,), center=(1, 7), ltr_cols=(3, 4))

# 7 Data & feedback
dd = [
 (T('Public aggregate admissions data', 'بيانات قبول إجمالية منشورة'), T('None found within the inventoried admissions scope (no application volumes, processing times or conversion figures).', 'لم يُعثر على شيء ضمن نطاق القبول المجرود (لا أعداد طلبات ولا مدد معالجة ولا نسب تحويل).'), T('No quantitative analysis is presented; no rates were calculated.', 'لا يُقدَّم تحليل كمي؛ ولم تُحسب أي نسب.')),
 (T('Applicant surveys / feedback summaries', 'استبيانات المتقدمين / ملخصات الآراء'), T('None published within scope; none supplied by the commissioning party.', 'لا يوجد منشور ضمن النطاق؛ ولم يُقدَّم شيء من صاحب التكليف.'), T('Experience findings are content-based observations of the information journey, not reported applicant experience.', 'ملاحظات التجربة مستمدة من مراجعة المحتوى لرحلة المعلومات، لا من تجربة متقدمين فعلية.')),
 (T('Information-journey review (content-based)', 'مراجعة رحلة المعلومات (مستندة إلى المحتوى)'), T('Four journeys traced: choose a programme → understand requirements → find fees and dates → find the application and follow-up channel; each stops at the portal boundary.', 'تُتبعت أربع رحلات: اختيار البرنامج ← فهم المتطلبات ← معرفة الرسوم والمواعيد ← الوصول إلى قناة التقديم والمتابعة؛ وتتوقف كل رحلة عند حدود البوابة.'), T('See report Section 6.', 'انظر القسم 6 من التقرير.')),
 (T('Proposed future design (not executed)', 'تصميم مقترح مستقبلًا (لم يُنفذ)'), T('(1) Analytics: page views, exits and search terms for admissions pages per cycle; (2) Enquiry log: categorise Admission Query topics monthly; (3) 6-item post-application survey (findability of requirements, fees, dates, channel; clarity; confidence), anonymous; (4) Review results before each cycle.', '(1) التحليلات: مشاهدات صفحات القبول ونقاط الخروج ومصطلحات البحث لكل دورة؛ (2) سجل الاستفسارات: تصنيف موضوعات نموذج الاستفسار شهريًا؛ (3) استبيان من 6 بنود بعد التقديم (سهولة العثور على المتطلبات والرسوم والمواعيد والقناة؛ الوضوح؛ الثقة)، دون هوية؛ (4) مراجعة النتائج قبل كل دورة.'), T('Requires separate authorisation, privacy review and owner approval.', 'يتطلب تكليفًا مستقلًا ومراجعة خصوصية واعتماد المالك.')),
]
sheet('Data & Feedback', 'البيانات والآراء', 'Data and applicant feedback', 'البيانات وآراء المتقدمين', '', '', [('Source / item', 'المصدر / البند'), ('Finding', 'النتيجة'), ('Limitation / implication', 'القيد / الأثر')], dd, [34, 80, 60])

# 8 Decisions
d_rows = [[d['did'], T(d['t_en'], d['t_ar']), T(d['q_en'], d['q_ar']), T(d['alt_en'], d['alt_ar']), ', '.join(d['fids']) if d['did'] != 'D-09' else T(f"{len(d['fids'])} editorial findings + {len(D['edits'])} annex corrections", f"{len(d['fids'])} ملاحظة تحريرية + {len(D['edits'])} تصحيحًا في المرفق"),
           T(d['owner_en'], d['owner_ar']), T('Pending — not approved', 'معلّق — غير معتمد'), '', '', ''] for d in D['decisions']]
sheet('Decisions Required', 'القرارات المطلوبة', 'Institutional decisions required', 'القرارات المؤسسية المطلوبة', 'Approval fields are intentionally blank. No value in this register is an approved policy.',
      'حقول الاعتماد فارغة عمدًا. لا تمثل أي قيمة في هذا السجل سياسة معتمدة.',
      [('ID', 'المعرف'), ('Subject', 'الموضوع'), ('Institutional question', 'السؤال المؤسسي'), ('Options', 'البدائل'), ('Linked findings', 'الملاحظات المرتبطة'), ('Proposed owner', 'الجهة المقترحة'),
       ('Status', 'الحالة'), ('Approved value / rule', 'القيمة / القاعدة المعتمدة'), ('Effective date', 'تاريخ السريان'), ('Approval reference', 'مرجع الاعتماد')],
      d_rows, [8, 30, 60, 46, 40, 30, 16, 24, 14, 18], center=(1, 7))

# 9 Follow-up plan
p_rows = [[p['aid'], T(p['en'], p['ar']), p['prio'], p['did'], p['owner'], T(p['kpi_en'], p['kpi_ar']), T('Not started — awaiting approval', 'لم يبدأ — بانتظار الاعتماد'), T('Re-read affected pages after publication and record evidence', 'إعادة قراءة الصفحات المتأثرة بعد النشر وتوثيق الدليل'), ''] for p in D['plan']]
sheet('Follow-up Plan', 'خطة المتابعة', 'Proposed follow-up plan', 'خطة المتابعة المقترحة', 'Proposed sequencing only; no timeline is committed on behalf of management.', 'تسلسل مقترح فقط؛ ولا يُلتزم بجدول زمني نيابة عن الإدارة.',
      [('Action', 'الإجراء'), ('Description', 'الوصف'), ('Priority', 'الأولوية'), ('Decision', 'القرار'), ('Proposed owner', 'الجهة المقترحة'), ('Indicator', 'المؤشر'), ('Implementation status', 'حالة التنفيذ'), ('Verification method', 'طريقة التحقق'), ('Verified on', 'تاريخ التحقق')],
      p_rows, [8, 60, 8, 9, 30, 34, 22, 34, 12], center=(1, 3, 4), prio_col=3)

# 10 Quality control
Q = D.get('qc', [])
sheet('Quality Control', 'ضبط الجودة', 'Quality control of outputs', 'ضبط جودة المخرجات', '', '', [('Check', 'الفحص'), ('Result', 'النتيجة'), ('Detail', 'التفاصيل')],
      [[T(q['en'], q['ar']), q['result'], q['detail']] for q in Q], [60, 12, 80], center=(2,))

# Annex sheets
sheet('Strengths', 'نقاط القوة', 'Proven strengths', 'نقاط القوة المثبتة', 'Each strength is tied to verbatim published evidence re-read on 21 Oct 2026.', 'كل نقطة قوة مرتبطة بدليل منشور حرفي أُعيدت قراءته في 21 أكتوبر 2026.',
      [('ID', 'المعرف'), ('Page', 'الصفحة'), ('URL', 'الرابط'), ('Evidence (verbatim)', 'الدليل (حرفيًا)'), ('Strength', 'نقطة القوة'), ('Verified', 'مُتحقق')],
      [[s['sid'], T(s['page_en'], s['page_ar']), s['url'], s['excerpt'], T(s['en'], s['ar']), 'Yes' if s['verified'] else 'No'] for s in D['strengths']], [8, 26, 36, 50, 60, 9], link_cols=(3,), center=(1, 6), ltr_cols=(3, 4))
sheet('Editorial Corrections', 'التصحيحات التحريرية', 'Ready editorial corrections (in-scope extract of the remediation package)', 'التصحيحات التحريرية الجاهزة (المستخلص ضمن النطاق من حزمة المعالجة)',
      'None changes a value, requirement or fee. Implement only after approval (D-09).', 'لا يغير أيٌّ منها قيمة أو شرطًا أو رسمًا. تُنفذ بعد الاعتماد فقط (D-09).',
      [('Change ID', 'المعرف'), ('URL', 'الرابط'), ('Current text (verbatim)', 'النص الحالي (حرفيًا)'), ('Proposed text', 'النص المقترح'), ('Verification', 'التحقق')],
      [[e['cid'], e['url'], e['before'], e['after'], T('Verified in published content' if e['verified'] else 'Needs manual check (dynamic portal label)', 'مُتحقق منه في المحتوى المنشور' if e['verified'] else 'يحتاج تحققًا يدويًا (تسمية ديناميكية في البوابة)')] for e in D['edits']],
      [10, 46, 40, 40, 30], link_cols=(2,), center=(1,), ltr_cols=(2, 3, 4))
sheet('Amendments 26 Oct 2025', 'تعديلات 26 أكتوبر 2025', 'Status of the amendments requested on 26 October 2025', 'حالة التعديلات المطلوبة في 26 أكتوبر 2025', 'Live check on 21 Oct 2026.', 'تحقق مباشر في 21 أكتوبر 2026.',
      [('ID', 'المعرف'), ('Requested amendment', 'التعديل المطلوب'), ('Status', 'الحالة'), ('Published evidence', 'الدليل المنشور'), ('Decision', 'القرار')],
      [[a['aid'], T(a['req_en'], a['req_ar']), T(a['st_en'], a['st_ar']), T(a['ev_en'], a['ev_ar']), a['did']] for a in D['amend']], [7, 50, 18, 80, 9], center=(1, 3, 5))
wb.properties.creator = 'Nader Alfarhan'; wb.properties.title = T('PMU Admissions Office — Content Review Register', 'سجل مراجعة محتوى مكتب القبول')
wb.save(OUT); print('saved', OUT, len(wb.sheetnames), wb.sheetnames)
