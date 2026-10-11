import openpyxl,collections
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
NAVY='1F3864';ACC='2E75B6';LIGHT='EAF1FB';GREY='595959'
thin=Side(style='thin',color='BFBFBF');B=Border(top=thin,bottom=thin,left=thin,right=thin)
F=lambda **k:Font(name='Calibri',**k)
def title(ws,t,sub,ncol):
    ws.insert_rows(1,3)
    ws['A1']=t;ws['A1'].font=F(size=16,bold=True,color=NAVY)
    ws['A2']=sub;ws['A2'].font=F(size=10,italic=True,color=GREY)
    for c in range(1,ncol+1): ws.cell(3,c).border=Border(bottom=Side(style='medium',color=NAVY))
    ws.row_dimensions[1].height=26
def body(ws,hr,widths,name,center=()):
    n=ws.max_column;last=ws.max_row
    for c in range(1,n+1):
        h=ws.cell(hr,c);h.font=F(bold=True,color='FFFFFF',size=10);h.fill=PatternFill('solid',fgColor=NAVY)
        h.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center');h.border=B
        ws.column_dimensions[L(c)].width=widths[c-1]
    ws.row_dimensions[hr].height=30
    for r in range(hr+1,last+1):
        for c in range(1,n+1):
            x=ws.cell(r,c);x.font=F(size=10);x.border=B
            x.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if c in center else 'left')
            if (r-hr)%2==0: x.fill=PatternFill('solid',fgColor=LIGHT)
    ws.freeze_panes=ws.cell(hr+1,1);ws.auto_filter.ref=f"A{hr}:{L(n)}{last}";ws.print_title_rows=f'{hr}:{hr}'
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_options.horizontalCentered=True
    ws.page_margins.left=ws.page_margins.right=0.4;ws.page_margins.top=ws.page_margins.bottom=0.5
    ws.oddFooter.center.text='Page &P of &N';ws.oddFooter.center.size=8
    ws.oddHeader.right.text=name;ws.oddHeader.right.size=8;ws.sheet_view.showGridLines=False
def prio(ws,col,hr):
    rng=f'{col}{hr+1}:{col}{ws.max_row}'
    for v,fg,bg in [('P1','9C0006','FFC7CE'),('P2','7F4F00','FFEB9C'),('P3','1E4620','C6EFCE')]:
        ws.conditional_formatting.add(rng,FormulaRule(formula=[f'{col}{hr+1}="{v}"'],font=Font(bold=True,color=fg),fill=PatternFill('solid',fgColor=bg)))
