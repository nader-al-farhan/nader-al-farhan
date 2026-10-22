# Master Directive workbooks. Usage: b360m_xlsx.py master.json out.xlsx en|ar L1|L2
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
D = json.load(open(sys.argv[1])); OUT = sys.argv[2]; AR = sys.argv[3] == 'ar'; LEVEL = sys.argv[4]
def T(en, ar): return ar if AR else en
NAVY = '1F3864'; GREY = '595959'
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
def F(**k): k.setdefault('name', 'Arial'); k.setdefault('size', 9); return Font(**k)
wb = Workbook(); first = [True]
PCOL = [('P0', '9C0006', 'FFC7CE'), ('P1', '7F4F00', 'FFEB9C'), ('P2', '1F3864', 'DDEBF7'), ('P3', '375623', 'E2EFDA'),
        ('PASS', '375623', 'E2EFDA'), ('FAIL', '9C0006', 'FFC7CE'), ('NOT VERIFIED', '7F4F00', 'FFEB9C'),
        ('NOT_PUBLISHED', '9C0006', 'FFC7CE'), ('PARTIALLY_PUBLISHED', '7F4F00', 'FFEB9C'), ('PUBLISHED', '375623', 'E2EFDA'),
        ('OPEN CONFLICT', '9C0006', 'FFC7CE'), ('ESCALATION REQUIRED', '7F4F00', 'FFEB9C'), ('BLOCKED', '7F4F00', 'FFEB9C'), ('VERIFIED', '375623', 'E2EFDA')]
def sheet(ten, tar, hen, har, sen, sar, cols, rows, widths, links=(), center=(), ltr=(), colour=()):
    ws = wb.active if first[0] else wb.create_sheet(); first[0] = False
    ws.title = T(ten, tar)[:31]; ws.sheet_view.rightToLeft = AR
    ws['A1'] = T(hen, har); ws['A1'].font = F(size=14, bold=True, color=NAVY)
    ws['A2'] = T(sen, sar); ws['A2'].font = F(size=9, italic=True, color=GREY)
    for j, c in enumerate(cols, 1):
        x = ws.cell(4, j, T(c[0], c[1])); x.font = F(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor=NAVY); x.border = B
        x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center', readingOrder=2 if AR else 1)
    for i, row in enumerate(rows, 5):
        for j, v in enumerate(row, 1):
            x = ws.cell(i, j, v); x.border = B; x.font = F()
            l = j in ltr
            x.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if j in center else ('left' if (l or not AR) else 'right'), readingOrder=1 if (l or not AR) else 2)
            if j in links and isinstance(v, str) and v.startswith('http'): x.hyperlink = v; x.font = F(color='0563C1', underline='single')
            if i % 2 == 0: x.fill = PatternFill('solid', fgColor='F7F9FC')
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 34
    if rows: ws.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{4 + len(rows)}'
    ws.freeze_panes = 'B5' if len(cols) > 3 else 'A5'
    for col in colour:
        L = get_column_letter(col)
        for v, fg, bg in PCOL:
            ws.conditional_formatting.add(f'{L}5:{L}{4 + max(1, len(rows))}', FormulaRule(formula=[f'ISNUMBER(SEARCH("{v}",{L}5))'], font=Font(bold=True, color=fg), fill=PatternFill('solid', fgColor=bg), stopIfTrue=True))
    ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = '4:4'
    return ws

R = D['rows']; S = D['stats']; C = D['coverage']
SUB = T('PMU-ADM-360-MASTER-20261022 · Prepared for Management Review · Not an approval, publication confirmation or closure certificate',
        'PMU-ADM-360-MASTER-20261022 · معد للمراجعة الإدارية · لا يمثل اعتمادًا أو تأكيد نشر أو شهادة إغلاق')

def decisions_sheet(title_en='Decision & Approval Matrix', title_ar='مصفوفة القرارات والاعتماد'):
    rows = [[d['did'], T(d['t_en'], d['t_ar']), T(d['q_en'], d['q_ar']), f"P0 {d['nP0']} · P1 {d['nP1']}", ', '.join(d['fids']) if d['did'] != 'D-09' else T(f"{len(d['fids'])} editorial findings + {len(D['edits'])} corrections", f"{len(d['fids'])} ملاحظة تحريرية + {len(D['edits'])} تصحيحًا"),
             T(d['risk_en'], d['risk_ar']), T(d['alt_en'], d['alt_ar']), T(d['rec_en'], d['rec_ar']), T(d['ref_en'], d['ref_ar']), T(d['auth'], d['auth_ar']), T(d['role'], d['role_ar']), T(d['impl'], d['impl_ar']), T(d['ver'], d['ver_ar']), T(d['status'], 'معلّق — غير معتمد'), '', d['follow']] for d in D['decisions']]
    sheet(title_en, title_ar, 'Decision & Approval Matrix', 'مصفوفة القرارات والاعتماد', SUB, SUB,
          [('ID', 'المعرف'), ('Matter to decide', 'المسألة'), ('Question & interpretation', 'السؤال والتفسير'), ('Linked severity', 'الخطورة المرتبطة'), ('Evidence (findings)', 'الأدلة (الملاحظات)'),
           ('Risk if unresolved', 'المخاطر عند الاستمرار'), ('Options', 'الخيارات'), ('Professional recommendation', 'التوصية المهنية'), ('Reference requiring approval', 'المرجع المطلوب اعتماده'),
           ('Decision authority', 'صاحب الصلاحية'), ("Mr. Radhi's role", 'دور الأستاذ راضي'), ('Implementation', 'جهة التنفيذ'), ('Publication verification', 'التحقق من النشر'),
           ('Status', 'الحالة'), ('Approval reference / date', 'مرجع الاعتماد / التاريخ'), ('Follow-up date', 'تاريخ المتابعة')],
          rows, [7, 26, 44, 12, 28, 36, 36, 40, 34, 26, 18, 24, 20, 16, 18, 18], center=(1, 4, 11, 14), ltr=(5,))

def roadmap_sheet(title_en='Corrective Action Roadmap', title_ar='خارطة الإجراءات التصحيحية'):
    rows = [[p['aid'], p['prio'], p['problem'], p['outcome'], p['action'], p['dec_pending'], p['owner'], p['approval'], p['target'], p['status'], p['evidence'], p['reverify'], T(p['kpi_en'], p['kpi_ar'])] for p in D['plan']]
    sheet(title_en, title_ar, 'Corrective Action Roadmap / Tracker', 'خارطة الإجراءات التصحيحية ومتابعتها',
          T('Lifecycle after authorisation: APPROVE → IMPLEMENT → SAVE → REOPEN → READ BACK → COMPARE → VERIFY → CLOSE. Target dates are proposals, not commitments.', 'دورة ما بعد التفويض: اعتماد ← تنفيذ ← حفظ ← إعادة فتح ← إعادة قراءة ← مقارنة ← تحقق ← إغلاق. المواعيد مقترحة وليست التزامًا.'), '',
          [('Action', 'الإجراء'), ('Priority', 'الأولوية'), ('1. Proven problem', '1. المشكلة المثبتة'), ('2. Required outcome', '2. النتيجة المطلوبة'), ('3. Proposed text / action', '3. النص أو الإجراء المقترح'),
           ('4. Decision pending', '4. القرار المعلّق'), ('5. Implementation owner', '5. مسؤول التنفيذ'), ('6. Prior approval', '6. الاعتماد السابق'), ('7. Proposed target', '7. الموعد المقترح'),
           ('8. Status', '8. حالة الإنجاز'), ('9. Implementation & publication evidence', '9. دليل التنفيذ والنشر'), ('10. Re-verification result', '10. نتيجة إعادة التحقق'), ('Indicator', 'المؤشر')],
          rows, [8, 8, 40, 32, 38, 10, 28, 28, 26, 14, 26, 22, 30], center=(1, 2, 6, 10), colour=(2,))

def amend_sheet(title_en='Historical Amendments', title_ar='التعديلات السابقة'):
    rows = [[a['hid'], T(a['req_en'], a['req_ar']), a['code'], T(a['ev_en'], a['ev_ar']), a['did'], a['checked'], T(a['basis_en'], a['basis_ar'])] for a in D['amend']]
    sheet(title_en, title_ar, 'Historical Amendment Verification (requests of 26 Oct 2025)', 'التحقق من التعديلات السابقة (طلبات 26 أكتوبر 2025)',
          T('Statuses: PUBLISHED · PARTIALLY_PUBLISHED · NOT_PUBLISHED · SUPERSEDED · CANNOT_VERIFY · NOT_APPLICABLE', 'الحالات: PUBLISHED · PARTIALLY_PUBLISHED · NOT_PUBLISHED · SUPERSEDED · CANNOT_VERIFY · NOT_APPLICABLE'), '',
          [('ID', 'المعرف'), ('Requested amendment', 'التعديل المطلوب'), ('Status', 'الحالة'), ('Current published evidence', 'الدليل المنشور الحالي'), ('Decision', 'القرار'), ('Checked', 'تاريخ التحقق'), ('Basis', 'الأساس')],
          rows, [7, 46, 22, 70, 9, 12, 46], center=(1, 3, 5, 6), colour=(3,))

def priority_sheet():
    rows = [[r['fid'], r['plevel'], T(r['title_en'], r['title_ar']), T(r['page_en'], r['page_ar']), T(r['cat_en'], r['cat_ar']), T(r['impact_en'], r['impact_ar']), r['res_status'], ', '.join(r['dids']), r['owner_code'], r['verified_on']]
            for r in R if r['plevel'] in ('P0', 'P1')]
    sheet('Priority Findings', 'الملاحظات ذات الأولوية', 'Priority findings (P0–P1)', 'الملاحظات ذات الأولوية (P0–P1)', SUB, SUB,
          [('ID', 'المعرف'), ('P', 'الأولوية'), ('Finding', 'الملاحظة'), ('Page', 'الصفحة'), ('Classification', 'التصنيف'), ('Impact', 'الأثر'), ('Resolution status', 'حالة الحسم'), ('Decision', 'القرار'), ('Owner', 'المسؤول'), ('Last verified', 'آخر تحقق')],
          rows, [8, 6, 44, 26, 26, 18, 22, 10, 26, 12], center=(1, 2, 7, 8, 10), colour=(2, 7))

RES = [(T('Open conflicts awaiting institutional confirmation', 'تعارضات مفتوحة بانتظار تأكيد مؤسسي'), S['res'].get('OPEN CONFLICT', 0), T('Applicants may receive contradictory requirements or fees', 'قد يتلقى المتقدم متطلبات أو رسومًا متناقضة'), 'D-01 … D-05'),
       (T('Items blocked from verification (F-033, F-131)', 'بنود تعذر التحقق منها (F-033، F-131)'), S['vstatus'].get('BLOCKED', 0), T('Possible spelling errors / failing http link remain unconfirmed', 'تبقى أخطاء إملائية محتملة ورابط http محتمل التعطل دون تأكيد'), T('Manual check from a public browser', 'فحص يدوي من متصفح عام')),
       (T('No approved policy references available to the review', 'لا مراجع سياسات معتمدة متاحة للمراجعة'), '—', T('Values cannot be confirmed as correct, only as conflicting', 'لا يمكن تأكيد صحة القيم، بل تعارضها فقط'), 'D-01 … D-05'),
       (T('No applicant data or feedback', 'لا بيانات أو آراء للمتقدمين'), '—', T('Experience findings are content-based', 'ملاحظات التجربة مستندة إلى المحتوى'), 'A-13'),
       (T('Operational application of procedures not verified', 'لم يُتحقق من التطبيق التشغيلي للإجراءات'), '—', T('Report does not state how procedures are applied', 'لا يحدد التقرير كيفية التطبيق'), 'D-01 … D-04'),
       (T('No correction published yet', 'لم يُنشر أي تصحيح بعد'), S['open'], T('All P0 risks remain live on 22 Oct 2026', 'جميع مخاطر P0 قائمة في 22 أكتوبر 2026'), 'A-01 … A-13'),
       (T('Exclusions by mandate (portal login, forms, security/performance testing, applicant records)', 'استثناءات بحكم التكليف (دخول البوابة، النماذج، اختبارات الأمن والأداء، سجلات المتقدمين)'), '—', T('Not assessed', 'غير مقيّمة'), '—')]
def residual_sheet():
    sheet('Gaps & Residual Risks', 'الفجوات والمخاطر المتبقية', 'Gaps, exceptions and residual risks', 'الفجوات والاستثناءات والمخاطر المتبقية', SUB, SUB,
          [('Gap / exception', 'الفجوة / الاستثناء'), ('Count', 'العدد'), ('Residual risk', 'المخاطر المتبقية'), ('Route', 'المسار')], [list(x) for x in RES], [60, 10, 60, 24], center=(2,))

if LEVEL == 'L1':
    decisions_sheet(); roadmap_sheet(); priority_sheet(); amend_sheet(); residual_sheet()
else:
    readme = [(T('Order', 'الأمر'), 'PMU-ADM-360-MASTER-20261022 · v3.0'), (T('Prepared for', 'معد لـ'), 'Mr. Radhi H. Alkhaldi — Associate Director of Enrollment'),
              (T('Prepared by', 'إعداد'), 'Nader Al Farhan — Administrative Assistant | Student Affairs'), (T('Root source', 'المصدر الجذري'), 'https://www.pmu.edu.sa/admission/admission'),
              (T('Verification dates', 'تواريخ التحقق'), D['obs_at']), (T('Records', 'السجلات'), f"{S['n']} findings (F) · {len(D['amend'])} historical requests (H) · {len(D['verified_correct'])} verified-correct records (V)"),
              ('P0 / P1 / P2 / P3', f"{S['plevel'].get('P0')} / {S['plevel'].get('P1')} / {S['plevel'].get('P2')} / {S['plevel'].get('P3')}"),
              (T('Verification status', 'حالة التحقق'), ' · '.join(f'{k} {v}' for k, v in S['vstatus'].items())),
              (T('Package status', 'حالة الحزمة'), 'READY FOR MANAGEMENT REVIEW'), (T('Closure status', 'حالة الإغلاق'), 'NOT READY FOR VERIFIED CLOSURE'),
              (T('Level', 'المستوى'), T('Level 2 — evidence and internal follow-up pack', 'المستوى الثاني — حزمة الأدلة والمتابعة الداخلية'))]
    sheet('README', 'اقرأني', 'PMU Admissions Audit 360 — Evidence & Follow-up Pack', 'تدقيق القبول 360 — حزمة الأدلة والمتابعة', SUB, SUB, [('Item', 'البند'), ('Value', 'القيمة')], readme, [30, 110])
    sheet('Source Register', 'سجل المصادر', 'Project source register', 'سجل مصادر المشروع', '', '',
          [('ID', 'المعرف'), ('Source', 'المصدر'), ('Location', 'الموقع'), ('Date / version', 'التاريخ / النسخة'), ('Reference authority', 'السلطة المرجعية'), ('Completeness', 'الاكتمال'), ('Access', 'حالة الوصول')],
          [list(s) for s in D['sources']], [9, 60, 26, 22, 36, 44, 18], center=(1, 7))
    mf = []
    for r in R:
        mf.append([r['fid'], T(r['page_en'], r['page_ar']), r['url'] if r['url'].startswith('http') else '', r['artifact'], r['lang_res'], r['current_text'], T(r['cat_en'], r['cat_ar']), T(r['nature_en'], r['nature_ar']),
                   T(r['rv_en'], r['rv_ar']), T(r['auth_ref'], r['auth_ref_ar']), r['auth_ver'], r['shot_ref'], T(r['impact_en'], r['impact_ar']), r['plevel'], T(r['rationale_en'], r['rationale_ar']),
                   T(r['description_en'], r['description']), T(r['recommendation_en'], r.get('recommendation_ar') or r['recommendation']), ', '.join(r['dids']), r['owner_code'],
                   r['vstatus'], r['verified_on'], r['res_status'], r['pub_status'], r['prior'], r['followup_en'], ', '.join(r['eids'])])
    for a in D['amend']:
        mf.append([a['hid'], T('Correspondence of 26 Oct 2025', 'مراسلة 26 أكتوبر 2025'), '', 'Email request', 'EN', T(a['req_en'], a['req_ar']), T(D['cat']['PRIOR_REQUEST_OPEN'][0], D['cat']['PRIOR_REQUEST_OPEN'][1]), a['code'], a['code'],
                   T(a['basis_en'], a['basis_ar']), '2025-10-26', '', T('Applicant · Academic/Financial', 'المتقدم · أكاديمي/مالي'), '', '', T(a['ev_en'], a['ev_ar']), '', a['did'], 'Admissions Office', 'VERIFIED', a['checked'], 'OPEN', 'Not published in full', '', '', ''])
    for v in D['verified_correct']:
        mf.append([v['vid'], T(v['page_en'], v['page_ar']), v['url'], 'HTML page', 'EN' if v['vid'] != 'V-10' else 'EN + AR', v['excerpt'], T(D['cat']['VERIFIED_CORRECT'][0], D['cat']['VERIFIED_CORRECT'][1]), T('Strength', 'نقطة قوة'), 'VERIFIED',
                   T('Published content (consistent across checked pages)', 'محتوى منشور (متسق في الصفحات المفحوصة)'), '2026-10-21/22', '', '', '', '', T(v['en'], v['ar']), T('Retain; use as single reference where applicable', 'الإبقاء عليه واستخدامه مرجعًا موحدًا عند الانطباق'), '', '', 'VERIFIED', '2026-10-21', 'NOT_APPLICABLE', 'Published', '', '', ''])
    sheet('Master Findings Register', 'السجل الرئيس للملاحظات', 'Master Findings Register (Single Source of Truth)', 'السجل الرئيس الموحد للملاحظات', SUB, SUB,
          [('Finding ID', 'المعرف'), ('Scope & page title', 'النطاق وعنوان الصفحة'), ('URL', 'الرابط'), ('Artifact type', 'نوع المورد'), ('Language', 'اللغة'), ('Observed claim or issue (verbatim)', 'الادعاء أو المسألة المرصودة (حرفيًا)'),
           ('Issue classification', 'تصنيف المسألة'), ('Nature', 'الطبيعة'), ('Current observed state', 'الحالة المرصودة حاليًا'), ('Authoritative reference', 'المرجع الحاكم'), ('Reference date & version', 'تاريخ المرجع ونسخته'),
           ('Screenshot / PDF page', 'اللقطة / صفحة PDF'), ('Impact', 'الأثر'), ('Severity', 'الخطورة'), ('Severity rationale', 'مبرر الخطورة'), ('Finding', 'الملاحظة'), ('Proposed correction', 'التصحيح المقترح'),
           ('Decision required', 'القرار المطلوب'), ('Owner', 'المسؤول'), ('Verification status', 'حالة التحقق'), ('Verified on', 'تاريخ التحقق'), ('Resolution status', 'حالة الحسم'), ('Publication status', 'حالة النشر'),
           ('Prior request', 'طلب سابق'), ('Follow-up notes', 'ملاحظات المتابعة'), ('Evidence IDs', 'معرفات الأدلة')],
          mf, [8, 24, 30, 12, 8, 44, 24, 14, 22, 36, 22, 20, 18, 7, 30, 50, 46, 9, 24, 12, 11, 20, 24, 8, 34, 14], links=(3,), center=(1, 5, 14, 18, 20, 21, 24), ltr=(3, 6, 12, 26), colour=(14, 20, 22))
    sheet('Evidence Register', 'سجل الأدلة', 'Evidence register', 'سجل الأدلة', T('Observation dates are separate from any source publication date.', 'تواريخ الرصد منفصلة عن تواريخ نشر المصادر.'), '',
          [('Evidence ID', 'معرف الدليل'), ('Finding', 'الملاحظة'), ('Source URL', 'رابط المصدر'), ('Excerpt (verbatim)', 'المقتطف (حرفيًا)'), ('Locator', 'الموضع'), ('First observed', 'أول رصد'), ('Last observed', 'آخر رصد'), ('Source date', 'تاريخ المصدر'), ('Highlighted screenshot', 'لقطة مظللة'), ('Full-page capture', 'لقطة كاملة')],
          [[e['eid'], e['fid'], e['url'], e['excerpt'], T(e['locator'], e['locator_ar']), e['first_observed'], e['observed'], e['pub_date'], e['shot'], e['page_shot']] for e in D['evidence']] +
          [[f"RC-{x['id']}", x['id'].rstrip('bc'), '', x['snippet'], T('Live re-confirmation 22 Oct 2026', 'إعادة تأكيد حي 22 أكتوبر 2026'), '', '2026-10-22', f"HTTP {x['http']} · matches {x['hits']}", '', ''] for x in D['recheck']],
          [12, 8, 40, 56, 30, 12, 12, 24, 18, 18], links=(3,), center=(1, 2, 6, 7), ltr=(3, 4, 9, 10))
    SN = D['sect']
    sheet('URL & Resource Inventory', 'جرد الروابط والموارد', 'URL & resource inventory', 'جرد الروابط والموارد', '', '',
          [('ID', 'المعرف'), ('Title', 'العنوان'), ('Type', 'النوع'), ('Language', 'اللغة'), ('Original URL', 'الرابط الأصلي'), ('Final URL', 'الرابط النهائي'), ('Discovered from', 'مسار الاكتشاف'), ('Section', 'القسم'),
           ('Inclusion', 'الإدراج'), ('Reason', 'السبب'), ('Accessed', 'تاريخ الوصول'), ('HTTP', 'HTTP'), ('Exam status', 'حالة الفحص'), ('Parts not examined', 'الأجزاء غير المفحوصة')],
          [[i['rid'], i['title'], i['type'], i['lang'], i['url'], i['final'], i['found'], T(SN[i['section']][0], SN[i['section']][1]), i['incl'], T(i['reason_en'], i['reason_ar']), i['at'], i['http'], i['status'], T(i['unchecked_en'], i['unchecked_ar'])] for i in D['inventory']],
          [10, 32, 16, 16, 40, 40, 30, 22, 16, 34, 11, 7, 11, 34], links=(5, 6), center=(1, 11, 12, 13), ltr=(2, 5, 6))
    import collections as _c
    cm = _c.defaultdict(lambda: _c.Counter())
    for i in D['inventory']: cm[i['section']][i['status']] += 1
    fc = _c.Counter(r['section'] for r in R)
    cov_rows = [[T(SN[k][0], SN[k][1]), sum(v.values()), v.get('FULL', 0), v.get('PARTIAL', 0), v.get('EXCLUDED', 0), 0, fc.get(k, 0)] for k, v in cm.items()]
    cov_rows.append([T('Total', 'المجموع'), len(D['inventory']), C['full'], C['partial'], C['excluded'], 0, S['n']])
    cov_rows.append([T(f"Coverage of eligible resources: full {C['pct_full']}% · any {C['pct_any']}% · links {C['links_unique']} unique ({len(C['broken'])} × 404) · PDFs {C['pdfs']} ({C['pdf_pages']} pp.)", f"تغطية الموارد المؤهلة: كاملة {C['pct_full']}% · أي فحص {C['pct_any']}% · روابط فريدة {C['links_unique']} ({len(C['broken'])} × 404) · PDF ‏{C['pdfs']} ({C['pdf_pages']} صفحة)"), '', '', '', '', '', ''])
    sheet('Scope Coverage Matrix', 'مصفوفة تغطية النطاق', 'Scope coverage matrix', 'مصفوفة تغطية النطاق', T('Coverage is proven from the current inventory, not inherited from earlier reports.', 'التغطية مثبتة من الجرد الحالي لا موروثة من تقارير سابقة.'), '',
          [('Section', 'القسم'), ('Resources', 'الموارد'), ('Fully examined', 'مفحوص كليًا'), ('Partially', 'جزئيًا'), ('Excluded', 'مستبعد'), ('Unreachable', 'تعذر الوصول'), ('Findings', 'الملاحظات')],
          cov_rows, [46, 12, 14, 12, 12, 12, 12], center=(2, 3, 4, 5, 6, 7))
    lp = [[l[0], l[1], 'https://www.pmu.edu.sa/apply' if l[0] != 'L-10' else 'https://pmu.edu.sa/admission/admissions_queries', l[2], l[3], T(l[4], l[5]), l[6]] for l in D['lang_pairs']]
    lp.append(['L-NOTE', T('Separate Arabic version', 'النسخة العربية المستقلة'), '', '', '', T(D['lang_note'][0], D['lang_note'][1]), ''])
    sheet('Arabic-English Parity', 'التوافق العربي الإنجليزي', 'Arabic / English parity register', 'سجل التوافق بين العربية والإنجليزية', '', '',
          [('Pair', 'الزوج'), ('Topic', 'الموضوع'), ('Source', 'المصدر'), ('English (verbatim)', 'الإنجليزية (حرفيًا)'), ('Arabic (verbatim)', 'العربية (حرفيًا)'), ('Result', 'النتيجة'), ('Finding', 'الملاحظة')],
          lp, [9, 24, 30, 44, 44, 50, 10], links=(3,), center=(1, 7), ltr=(3, 4))
    amend_sheet('Historical Amendment Verif.', 'التحقق من التعديلات السابقة')
    decisions_sheet('Decision Register', 'سجل القرارات')
    roadmap_sheet('Corrective Action Tracker', 'متابعة الإجراءات التصحيحية')
    sheet('QC & Acceptance Matrix', 'مصفوفة الجودة والقبول', 'Quality gates and acceptance matrix', 'معايير الجودة ومصفوفة القبول', T('Gate results: PASS / FAIL / NOT VERIFIED', 'نتائج المعايير: PASS / FAIL / NOT VERIFIED'), '',
          [('Gate / check', 'المعيار / الفحص'), ('Result', 'النتيجة'), ('Evidence', 'الدليل'), ('Residual note', 'الملاحظة المتبقية'), ('Action to satisfy', 'الإجراء المطلوب')],
          [[T(g['en'], g['ar']), g['result'], g['evidence'], g.get('residual', ''), g.get('action', '')] for g in D.get('gates', [])], [52, 14, 60, 50, 44], center=(2,), colour=(2,))
    sheet('Change Log', 'سجل التغييرات', 'Change log', 'سجل التغييرات', '', '', [('Version', 'النسخة'), ('Date', 'التاريخ'), ('Change', 'التغيير')], [list(c) for c in D['changelog']], [10, 12, 120], center=(1, 2))
    sheet('Editorial Corrections', 'التصحيحات التحريرية', 'Ready editorial corrections (D-09)', 'التصحيحات التحريرية الجاهزة (D-09)', T('None changes a value, requirement or fee.', 'لا يغيّر أيٌّ منها قيمة أو شرطًا أو رسمًا.'), '',
          [('Change ID', 'المعرف'), ('URL', 'الرابط'), ('Current text', 'النص الحالي'), ('Proposed text', 'النص المقترح'), ('Verification', 'التحقق')],
          [[e['cid'], e['url'], e['before'], e['after'], 'VERIFIED' if e['verified'] else 'BLOCKED (dynamic portal label)'] for e in D['edits']], [10, 46, 40, 40, 26], links=(2,), center=(1, 5), ltr=(2, 3, 4), colour=(5,))
    defs = [(k, T(v[0], v[1])) for k, v in D['cat'].items()] + [(k, T(v[0], v[1])) for k, v in D['pl'].items()] + [
        ('VERIFIED', T('The observed state was re-read and confirmed against retrievable evidence on the stated date.', 'أُعيدت قراءة الحالة المرصودة وتأكيدها بدليل قابل للاسترجاع في التاريخ المذكور.')),
        ('BLOCKED', T('Verification could not be completed with permitted means; reason recorded.', 'تعذر إكمال التحقق بالوسائل المسموح بها؛ والسبب مسجل.')), ('NOT_APPLICABLE', T('Verification not applicable.', 'التحقق غير منطبق.')),
        ('OPEN CONFLICT', T('Two published sources conflict and no governing reference resolves them.', 'تعارض مصدرين منشورين دون مرجع حاكم يحسم.')), ('ESCALATION REQUIRED', T('A P0/P1 matter needing a decision by a competent authority.', 'مسألة P0/P1 تحتاج قرارًا من صاحب صلاحية.')),
        ('OWNER_TO_CONFIRM', T('No evidence identifies the responsible unit.', 'لا دليل يحدد الجهة المسؤولة.')), ('DIRECT / REFER', T("Matters Mr. Radhi can direct within Admissions vs. matters to refer to the competent authority.", 'مسائل يوجه فيها الأستاذ راضي ضمن القبول مقابل مسائل تُحال لصاحب الصلاحية.')),
        ('READY FOR MANAGEMENT REVIEW', T('Package documented, decisions clear, gaps disclosed, gates met for handover.', 'حزمة موثقة وقرارات واضحة وفجوات معلنة ومعايير مستوفاة للتسليم.')),
        ('VERIFIED-OUTCOME-CLOSED', T('Only after decisions, implementation, publication, re-verification and accepted residual risk.', 'فقط بعد القرارات والتنفيذ والنشر وإعادة التحقق وقبول المخاطر المتبقية.')),
        ('NOT READY FOR VERIFIED CLOSURE', T('Any material issue, approval or unverified correction outstanding.', 'بقاء أي خلل جوهري أو موافقة أو تصحيح غير متحقق.'))]
    sheet('Definitions', 'التعريفات', 'Definitions', 'التعريفات', '', '', [('Term', 'المصطلح'), ('Definition', 'التعريف')], defs, [30, 110])
wb.properties.creator = 'Nader Al Farhan'; wb.properties.title = 'PMU Admissions Audit 360'
wb.save(OUT); print('saved', OUT, wb.sheetnames)
