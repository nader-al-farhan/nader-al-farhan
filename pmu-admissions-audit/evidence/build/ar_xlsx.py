import json,sys,openpyxl
sys.path.insert(0,'evidence/build');from xlstyle import *
from openpyxl.styles import Alignment
R=json.load(open('final_review/final_rows.json'));PM=json.load(open('final_review/pagemap.json'))
ORDER=['P00','P03','P04','P05','P06','P07','P08','P38','Shared menu','P09','P24','P25','P26','P27','P28','P29','P30','P31','P10','P11','P12','P13','P14','P15','P16','P17','P18','P19','P20','P21','P22','P23','P32','P33','P34','P35','P36','P37']
PR={'عالية':0,'متوسطة':1,'منخفضة':2}
R.sort(key=lambda a:(ORDER.index(a['page_id']),PR[a['priority']]))
pn=lambda i:'القائمة الجانبية لصفحات القبول' if i=='Shared menu' else PM[i]['ar']
wb=openpyxl.Workbook();ws=wb.active;ws.title='سجل الملاحظات'
ws.append(['#','الصفحة','الرابط','الموقع','النص الحالي (حرفيًا)','المحور','نوع الملاحظة','الوصف','التصنيف','الأولوية','التصحيح أو التوصية المقترحة','ملاحظة'])
for i,r in enumerate(R,1):
    ws.append([i,pn(r['page_id']),r['url'],r['location'],r['current_text'],r['axis'],r['type'],r['description'],r['classification'],r['priority'],r['recommendation'],r.get('note','')])
title(ws,'سجل الملاحظات الموحّد – صفحات مكتب القبول','المحتوى المنشور كما رُصد في 10 أكتوبر 2026 مع إعادة تحقق في 11 أكتوبر 2026',12)
body(ws,4,[5,22,30,24,48,14,20,48,13,10,48,30],'تقرير المراجعة الشاملة لصفحات مكتب القبول',center=(1,9,10))
for row in ws.iter_rows(min_row=5):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if c.column in (1,9,10) else ('left' if c.column in (3,5) else 'right'),readingOrder=2 if c.column not in (3,5) else 1)
for v,fg,bg in [('عالية','9C0006','FFC7CE'),('متوسطة','7F4F00','FFEB9C'),('منخفضة','1E4620','C6EFCE')]:
    ws.conditional_formatting.add(f'J5:J{ws.max_row}',FormulaRule(formula=[f'J5="{v}"'],font=Font(bold=True,color=fg),fill=PatternFill('solid',fgColor=bg)))
ws.sheet_view.rightToLeft=True
s=wb.create_sheet('الإحصاء');s.sheet_view.rightToLeft=True;s.sheet_view.showGridLines=False
s['A1']='الإحصاء الختامي';s['A1'].font=F(size=16,bold=True,color=NAVY)
s['A2']=f'العدد الإجمالي: {len(R)} ملاحظة (مطابق لعدد صفوف السجل). المجاميع أدناه محسوبة بمعادلات من ورقة السجل.';s['A2'].font=F(size=10,italic=True,color=GREY)
def tbl(r0,title_,labels,col):
    s.cell(r0,1,title_).font=F(size=12,bold=True,color=NAVY)
    hdr=['البند','عالية','متوسطة','منخفضة','المجموع']
    for i,h in enumerate(hdr):
        c=s.cell(r0+1,1+i,h);c.font=F(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=NAVY);c.alignment=Alignment(horizontal='center');c.border=B
    for k,lab in enumerate(labels):
        rr=r0+2+k;s.cell(rr,1,lab).border=B
        for j,p in enumerate(['عالية','متوسطة','منخفضة']):
            c=s.cell(rr,2+j,f"=COUNTIFS('سجل الملاحظات'!${col}$5:${col}$500,$A{rr},'سجل الملاحظات'!$J$5:$J$500,\"{p}\")");c.border=B;c.alignment=Alignment(horizontal='center')
        c=s.cell(rr,5,f'=SUM(B{rr}:D{rr})');c.border=B;c.font=F(bold=True);c.alignment=Alignment(horizontal='center')
    rr=r0+2+len(labels);s.cell(rr,1,'المجموع').font=F(bold=True);s.cell(rr,1).border=B
    for j,colL in enumerate('BCDE'):
        c=s.cell(rr,2+j,f'=SUM({colL}{r0+2}:{colL}{rr-1})');c.font=F(bold=True);c.border=B;c.alignment=Alignment(horizontal='center');c.fill=PatternFill('solid',fgColor=LIGHT)
    return rr+2
n=tbl(4,'حسب التصنيف',['تصحيح مطلوب','توصية تحسين','متابعة وتحقق'],'I')
tbl(n,'حسب المحور',['دقة المعلومات','لغوي ونحوي','الاتساق','تجربة المستخدم','وظيفي وتقني','تنظيمي'],'F')
s.column_dimensions['A'].width=28
for c in 'BCDE': s.column_dimensions[c].width=11
sc=wb.create_sheet('نطاق الرصد');sc.sheet_view.rightToLeft=True
sc.append(['المعرّف','الصفحة','الرابط','حالة الفحص'])
for i in sorted(k for k in PM if k.startswith('P') and k[1:].isdigit()):
    st='فُحص الرابط ووجهته فقط (بيانات شخصية)' if i=='P06' else ('خارج النطاق – صفحة عامة' if i in('P01','P02') else 'فُحصت سطرًا بسطر')
    sc.append([i,PM[i]['ar'] or PM[i]['title'],PM[i]['url'],st])
sc.append(['PDF','دليل المرشح لاختبار Aptis (2017)','https://www.pmu.edu.sa/Attachments/Admission/PDF/Appendix%203.3%20(Aptis%20Candidate%20Guide).pdf','فُحص'])
sc.append(['PDF','دليل التقديم الإلكتروني (2018)','https://www.pmu.edu.sa/Attachments/Admission/Images/PDF/Web-Admission-Guide-v2-2_20-12-18.pdf','فُحص'])
title(sc,'نطاق الرصد','الصفحات والملفات التي فُحصت وحالتها',4)
body(sc,4,[9,40,70,36],'تقرير المراجعة الشاملة لصفحات مكتب القبول',center=(1,))
wb.save('final_review/سجل_ملاحظات_صفحات_مكتب_القبول.xlsx');print('ok',len(R))
